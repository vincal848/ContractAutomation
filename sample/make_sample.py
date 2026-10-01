"""Builds sample/FinancialReport_sample.xlsx and sample/ContractTemplate.docx.

Run with `python sample/make_sample.py` from the repo root. Both outputs are
also committed, so this script only needs to be re-run if the sample data or
template layout changes.

python-docx has no bookmark API, so ContractTemplate.docx is built by
inserting raw w:bookmarkStart/w:bookmarkEnd XML around each placeholder run --
the same elements Word itself writes, just without going through Word.
"""
import os

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# Reconstructed from docs/img/result_image_sample.png / finautomation_results.pdf:
# Total Revenue and Net Profit are given directly there; Expenses = Revenue - Net Profit.
MONTHS = ["January", "February", "March", "April", "May", "June",
          "July", "August", "September", "October", "November", "December"]
REVENUES = [100000, 120000, 115000, 130000, 125000, 140000,
            150000, 145000, 155000, 160000, 170000, 180000]
EXPENSES = [60000, 70000, 65000, 80000, 75000, 85000,
            90000, 88000, 95000, 100000, 110000, 115000]
REGION = "Northeast"
CUSTOMER_NAME = "Sample Client"
COMPANY_NAME = "Sample Company"

BOOKMARKS = [
    "RegionBookmark",
    "MonthBookmark",
    "RevenueBookmark",
    "ExpensesBookmark",
    "NetProfitBookmark",
    "CustomerNameBookmark",
    "CompanyNameBookmark",
]


def build_workbook(path):
    wb = Workbook()

    ws_sources = wb.active
    ws_sources.title = "DataSources"
    ws_sources["A1"] = "Source Workbook Path"
    ws_sources["A1"].font = Font(bold=True)
    # Left empty below the header: this sample workbook ships with RawData
    # already filled in, to demonstrate CalculateMetrics / FormatSummaryReport
    # / CreateCharts / TransferToWordTemplate without needing real external
    # workbooks on disk. Add paths here (one per row, from row 2 down) to also
    # exercise AutomateFinancialReporting's consolidation step.

    ws_data = wb.create_sheet("RawData")
    headers = ["Region", "Month", "Revenue", "Expenses", "CustomerName", "CompanyName"]
    ws_data.append(headers)
    for cell in ws_data[1]:
        cell.font = Font(bold=True)
    for month, revenue, expense in zip(MONTHS, REVENUES, EXPENSES):
        ws_data.append([REGION, month, revenue, expense, CUSTOMER_NAME, COMPANY_NAME])

    ws_metrics = wb.create_sheet("CalculatedMetrics")
    ws_metrics.append(["Net Profit", "Revenue Growth (%)", "Profit Margin (%)", "Total Revenue"])
    for cell in ws_metrics[1]:
        cell.font = Font(bold=True)
    # Left otherwise empty: AutomateFinancialReporting (or CalculateMetrics
    # directly) fills this sheet in when the macro runs.

    ws_template = wb.create_sheet("TemplateInfo")
    ws_template["A1"] = "Template Path"
    ws_template["B1"] = "ContractTemplate.docx"
    ws_template["A2"] = "Output Folder"
    ws_template["B2"] = "Contracts"
    for cell in ("A1", "A2"):
        ws_template[cell].font = Font(bold=True)
    # B1/B2 are relative placeholders for this sample. Replace both with
    # absolute paths on your machine before running TransferToWordTemplate --
    # see README.md Setup.

    wb.save(path)


def _add_bookmark(paragraph, bookmark_id, bookmark_name, label_text):
    """Inserts a bookmarked run (label_text) into paragraph via raw XML,
    since python-docx exposes no bookmark API."""
    run = paragraph.add_run()

    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), bookmark_name)
    run._r.addprevious(start)

    text_run = paragraph.add_run(label_text)

    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    text_run._r.addnext(end)


def build_template(path):
    doc = Document()
    doc.add_heading("Contract", level=1)

    labels = {
        "RegionBookmark": "[Region]",
        "MonthBookmark": "[Month]",
        "RevenueBookmark": "[Revenue]",
        "ExpensesBookmark": "[Expenses]",
        "NetProfitBookmark": "[Net Profit]",
        "CustomerNameBookmark": "[Customer Name]",
        "CompanyNameBookmark": "[Company Name]",
    }

    # One paragraph per field, so each bookmark wraps a run on its own line.
    for i, bookmark_name in enumerate(BOOKMARKS, start=1):
        p = doc.add_paragraph()
        p.add_run(bookmark_name.replace("Bookmark", "") + ": ")
        _add_bookmark(p, i, bookmark_name, labels[bookmark_name])

    doc.save(path)


if __name__ == "__main__":
    build_workbook(os.path.join(ROOT, "sample", "FinancialReport_sample.xlsx"))
    build_template(os.path.join(ROOT, "sample", "ContractTemplate.docx"))
    print("Wrote sample/FinancialReport_sample.xlsx and sample/ContractTemplate.docx")
